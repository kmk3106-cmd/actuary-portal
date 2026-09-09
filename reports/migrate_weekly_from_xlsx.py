"""
계리결산팀 원본 주간업무 xlsx → 포탈 weekly_tasks 테이블 마이그.

- 입력 xlsx : --src (기본: Downloads/계리결산팀_주간업무_26년9월2주차.xlsx)
- 인별 시트 '주간업무_인별' 만 읽음 (다른 시트는 무시)
- 카테고리(B열) → 업무유형(Ⅰ~Ⅵ) 자동 매핑 (WORK_TYPE_MAP)
- 첫 마이그이므로 extension_type = '' (팀장이 나중에 지정)
- this_week_done / next_week_plan 도 비워둠

사용:
  python reports/migrate_weekly_from_xlsx.py --src <xlsx> --year 2026 --week 37 \
      --label "'26년 9월 2주차" --base 2026-09-07 [--dry-run]
"""
import argparse
import http.client
import json
import sys
import os
from datetime import datetime
from openpyxl import load_workbook


# 카테고리 → 업무유형 (Ⅰ~Ⅵ) 매핑
WORK_TYPE_MAP = [
    ('IFRS17 결산',         'Ⅰ'),
    ('결산',                'Ⅱ'),
    ('준비금',              'Ⅱ'),
    ('비금',                'Ⅱ'),
    ('계약',                'Ⅱ'),
    ('차세대',              'Ⅲ'),
    ('통합계리',            'Ⅳ'),
    ('시스템·로직',         'Ⅳ'),
    ('시스템',              'Ⅳ'),
    ('로직',                'Ⅳ'),
    ('모델',                'Ⅳ'),
    ('대내외',              'Ⅴ'),
    ('발송',                'Ⅴ'),
    ('감독',                'Ⅴ'),
    ('회계법인',            'Ⅴ'),
    ('계리법인',            'Ⅴ'),
    ('관리회계',            'Ⅵ'),
    ('사업계획',            'Ⅵ'),
    ('기타',                'Ⅵ'),
]

def infer_work_type(cat):
    c = (cat or '').strip()
    for kw, code in WORK_TYPE_MAP:
        if kw in c:
            return code
    return 'Ⅵ'  # 기본값


def _norm_date(v):
    if v is None or v == '':
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = str(v).strip()
    # 원본에 '계속', '상시', '수시' 같은 문자열도 있음 → 그대로
    return s


def _norm_progress(v):
    if v is None or v == '':
        return None
    try:
        f = float(v)
        return max(0.0, min(1.0, f))
    except Exception:
        return None


def parse_person_sheet(path):
    wb = load_workbook(path, data_only=True)
    ws = wb['주간업무_인별']
    # 헤더 :  A=팀원 B=구분 C=업무내용 D=이슈 E=시작일 F=완료 G=진척율 H=상태 I=팀원KEY
    # 팀원 A열은 병합, 구분 B열은 병합, 각 행별 데이터가 있음
    tasks = []
    for row in range(4, ws.max_row + 1):
        # I열(팀원KEY)이 채워져 있는 행만 유효 데이터
        member = ws.cell(row=row, column=9).value
        content = ws.cell(row=row, column=3).value
        if not member or not content:
            continue
        # ▶ 접두어 제거
        content_str = str(content).strip()
        if content_str.startswith('▶'):
            content_str = content_str.lstrip('▶').strip()
        # ' |     - X' 같은 노이즈는 그대로 유지 (원본 유지)
        task = {
            'member_name':      str(member).strip(),
            'category':         str(ws.cell(row=row, column=2).value or '').strip(),
            'extension_type':   '',
            'task_content':     content_str,
            'this_week_done':   '',
            'next_week_plan':   '',
            'issue_note':       str(ws.cell(row=row, column=4).value or '').strip(),
            'start_date':       _norm_date(ws.cell(row=row, column=5).value),
            'end_date':         _norm_date(ws.cell(row=row, column=6).value),
            'progress':         _norm_progress(ws.cell(row=row, column=7).value),
            'status':           str(ws.cell(row=row, column=8).value or '').strip(),
        }
        task['work_type']       = infer_work_type(task['category'])
        task['work_type_detail']= task['category']
        tasks.append(task)
    return tasks


def post_task(host, port, week_meta, task):
    conn = http.client.HTTPConnection(host, port, timeout=10)
    body = dict(week_meta)
    body.update(task)
    payload = json.dumps(body, ensure_ascii=False).encode('utf-8')
    conn.request('POST', '/tables/weekly_tasks', body=payload,
                 headers={'Content-Type': 'application/json; charset=utf-8', 'Content-Length': str(len(payload))})
    r = conn.getresponse()
    d = r.read().decode('utf-8', errors='replace')
    conn.close()
    return r.status, d


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=os.path.expandvars(r'%USERPROFILE%\Downloads\계리결산팀_주간업무_26년9월2주차.xlsx'))
    ap.add_argument('--year', type=int, required=True)
    ap.add_argument('--week', type=int, required=True)
    ap.add_argument('--label', required=True)
    ap.add_argument('--base', required=True, help='기준일 YYYY-MM-DD')
    ap.add_argument('--host', default='localhost')
    ap.add_argument('--port', type=int, default=8888)
    ap.add_argument('--dry-run', action='store_true', help='POST 없이 파싱 결과만 출력')
    ap.add_argument('--clear-existing', action='store_true', help='해당 주차 기존 rows 먼저 DELETE')
    args = ap.parse_args()

    if not os.path.exists(args.src):
        print(f'ERROR: xlsx 없음: {args.src}', file=sys.stderr); sys.exit(2)

    tasks = parse_person_sheet(args.src)
    print(f'파싱 완료 : {len(tasks)}건')
    # 요약
    by_member = {}
    for t in tasks:
        by_member.setdefault(t['member_name'], []).append(t)
    for m, ts in by_member.items():
        print(f'  {m} : {len(ts)}건 (유형 {sorted({t["work_type"] for t in ts})})')

    if args.dry_run:
        print('--dry-run · POST 안 함. 첫 3건:')
        for t in tasks[:3]:
            print('   ', json.dumps(t, ensure_ascii=False)[:180])
        return

    if args.clear_existing:
        # 해당 주차 rows 삭제
        conn = http.client.HTTPConnection(args.host, args.port, timeout=10)
        conn.request('GET', '/tables/weekly_tasks?limit=2000')
        r = conn.getresponse(); d = json.loads(r.read().decode('utf-8')); conn.close()
        exist = [x for x in (d.get('rows') or d.get('data') or [])
                 if int(x.get('year', 0)) == args.year and int(x.get('week_no', 0)) == args.week]
        print(f'기존 {len(exist)}건 삭제 중…')
        for x in exist:
            conn = http.client.HTTPConnection(args.host, args.port, timeout=10)
            conn.request('DELETE', '/tables/weekly_tasks/' + x['id'])
            conn.getresponse().read(); conn.close()

    week_meta = {
        'year': args.year, 'week_no': args.week, 'week_label': args.label,
        'base_date': args.base,
    }
    ok, fail = 0, 0
    for t in tasks:
        st, body = post_task(args.host, args.port, week_meta, t)
        if 200 <= st < 300:
            ok += 1
        else:
            fail += 1
            print(f'FAIL {st}: {body[:200]}')
    print(f'\n결과 : {ok}건 성공 · {fail}건 실패')


if __name__ == '__main__':
    main()
