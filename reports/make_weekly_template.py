"""
계리결산팀 주간업무 인별 템플릿 생성기.

- 참고 양식: C:\\Users\\USER\\Downloads\\계리결산팀_주간업무_26년9월2주차.xlsx
- 산출: 지정된 --out 경로 (.xlsx) — 3 시트
  1) 주간업무_인별  : 팀원별 블록(A merged) + 구분/업무내용/이슈/시작일/완료/진척율/상태 컬럼
  2) 요약           : COUNTIF/AVERAGEIFS 로 인별 건수·완료·진행중·미착수·평균진척율 자동 집계
  3) 통합_업무유형별 : 헤더만 있는 스켈레톤 (원하면 나중에 인별에서 복사·재분류)

- 팀장(김민국) 는 팀원 리스트에서 제외 (인별 실무 배정 대상 아님)
- 팀원 순서는 원본 양식 순서 유지 + 신규 팀원(이성원)은 마지막에 추가
- 팀원당 기본 5행 (--rows-per-person 로 조정 가능)
- 상태 셀 : 데이터 유효성(드롭다운) 미착수 / 진행중 / 완료
- 진척율 : 0.0% 서식 (0~1 저장 → 셀에 %로 표시)
- 시작일/완료(예정) : yyyy-mm-dd 서식
- 팀원KEY(I열) : 요약 시트 집계용 helper — 각 행의 팀원 이름을 그대로 넣음 (병합 X)
"""
import argparse
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles.differential import DifferentialStyle


# ── 팀원 리스트 (팀장 제외, 원본 양식 순서 유지) ─────────────────────
DEFAULT_MEMBERS = [
    '한인석', '고인수', '마혜원', '이상현', '이동민', '예대호',
    '김예은', '강세진', '오정택', '이용우', '김채린', '이성원',
]

# ── 스타일 상수 (원본 파일 톤앤매너: 남색 헤더 + 회색 소구분) ────────
TITLE_FONT   = Font(name='맑은 고딕', size=15, bold=True, color='FFFFFF')
SUB_FONT     = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
HEAD_FONT    = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
BODY_FONT    = Font(name='맑은 고딕', size=10)
NAME_FONT    = Font(name='맑은 고딕', size=11, bold=True, color='0F172A')
CAT_FONT     = Font(name='맑은 고딕', size=10, bold=True, color='334155')
FOOT_FONT    = Font(name='맑은 고딕', size=9, italic=True, color='6B7280')

TITLE_FILL   = PatternFill('solid', fgColor='1E3A5F')  # 남색
SUB_FILL     = PatternFill('solid', fgColor='334155')  # 진회색
HEAD_FILL    = PatternFill('solid', fgColor='475569')  # 회색 헤더
NAME_FILL    = PatternFill('solid', fgColor='EFF6FF')  # 옅은 파랑 (팀원 셀)
CAT_FILL     = PatternFill('solid', fgColor='F1F5F9')  # 옅은 회색 (구분 셀)
FOOT_FILL    = PatternFill('solid', fgColor='F9FAFB')

THIN = Side(border_style='thin', color='CBD5E1')
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT   = Alignment(horizontal='left',   vertical='center', wrap_text=True)


def build_person_sheet(ws, week_label, base_date, members, rows_per_person):
    ws.title = '주간업무_인별'
    # 열너비 (원본 참고)
    widths = {'A': 10, 'B': 18, 'C': 55, 'D': 26, 'E': 12, 'F': 12, 'G': 10, 'H': 10, 'I': 10}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # 1행: 제목
    ws.merge_cells('A1:H1')
    ws['A1'] = '계리결산팀 주간업무 현황'
    ws['A1'].font = TITLE_FONT
    ws['A1'].fill = TITLE_FILL
    ws['A1'].alignment = CENTER
    ws.row_dimensions[1].height = 28

    # 2행: 주차 · 기준일
    ws.merge_cells('A2:H2')
    ws['A2'] = f"{week_label}   |   기준일 : {base_date.strftime('%Y.%m.%d')}"
    ws['A2'].font = SUB_FONT
    ws['A2'].fill = SUB_FILL
    ws['A2'].alignment = CENTER
    ws.row_dimensions[2].height = 22

    # 3행: 컬럼 헤더
    headers = ['팀원', '구분', '업무내용', '이슈 · 비고', '시작일', '완료(예정)', '진척율', '상태', '팀원KEY']
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER_ALL
    ws.row_dimensions[3].height = 24
    # 팀원KEY 열은 helper — 열을 숨김 처리
    ws.column_dimensions['I'].hidden = True

    # ── 상태 드롭다운 (전 인별 블록에 적용) ───────────
    dv_status = DataValidation(
        type='list',
        formula1='"미착수,진행중,완료"',
        allow_blank=True,
        showDropDown=False,  # False = 드롭다운 화살표를 실제로 표시
    )
    dv_status.error = '미착수 / 진행중 / 완료 중 하나를 선택하세요.'
    dv_status.errorTitle = '상태 값 오류'
    ws.add_data_validation(dv_status)

    dv_pct = DataValidation(
        type='decimal',
        operator='between',
        formula1=0, formula2=1,
        allow_blank=True,
    )
    dv_pct.error = '0 ~ 1 사이 값 (0.5 = 50%)'
    dv_pct.errorTitle = '진척율 오류'
    ws.add_data_validation(dv_pct)

    # ── 팀원별 블록 ────────────────────────────────
    row = 4
    for name in members:
        top = row
        bot = row + rows_per_person - 1
        # A 열 팀원 이름 병합
        ws.merge_cells(start_row=top, start_column=1, end_row=bot, end_column=1)
        a = ws.cell(row=top, column=1, value=name)
        a.font = NAME_FONT
        a.fill = NAME_FILL
        a.alignment = CENTER
        # 각 행 스타일 + I 열에 팀원 KEY 채움
        for r in range(top, bot + 1):
            for col in range(1, 10):
                cell = ws.cell(row=r, column=col)
                cell.border = BORDER_ALL
                if col in (5, 6, 7, 8):  # 시작일·완료·진척율·상태
                    cell.alignment = CENTER
                else:
                    cell.alignment = LEFT
                cell.font = BODY_FONT
            # 팀원KEY (I열) — 각 행에 이름 저장 (요약 집계 helper)
            ws.cell(row=r, column=9, value=name).font = BODY_FONT
            # 서식
            ws.cell(row=r, column=5).number_format = 'yyyy-mm-dd'
            ws.cell(row=r, column=6).number_format = 'yyyy-mm-dd'
            ws.cell(row=r, column=7).number_format = '0.0%'
            ws.row_dimensions[r].height = 22
        # 데이터 유효성 범위 등록
        dv_status.add(f'H{top}:H{bot}')
        dv_pct.add(f'G{top}:G{bot}')
        row = bot + 1

    # 하단 주석 (병합)
    last = row
    ws.merge_cells(start_row=last, start_column=1, end_row=last, end_column=8)
    ws.cell(row=last, column=1, value=(
        '※ 상태 구분 : 완료(진척율 100%) / 진행중(0% 초과 ~ 100% 미만) / 미착수(0%)   |   '
        '진척율 미기재 항목은 상시·수시·착수 전 과제 (평균 진척율 집계에서 제외)'
    ))
    ws.cell(row=last, column=1).font = FOOT_FONT
    ws.cell(row=last, column=1).fill = FOOT_FILL
    ws.cell(row=last, column=1).alignment = LEFT

    # 상태별 조건부 서식 (완료=녹색, 진행중=파랑, 미착수=회색)
    data_range = f'H4:H{row - 1}'
    green_fill  = PatternFill('solid', fgColor='D1FAE5')
    green_font  = Font(name='맑은 고딕', size=10, bold=True, color='065F46')
    blue_fill   = PatternFill('solid', fgColor='DBEAFE')
    blue_font   = Font(name='맑은 고딕', size=10, bold=True, color='1E40AF')
    gray_fill   = PatternFill('solid', fgColor='F1F5F9')
    gray_font   = Font(name='맑은 고딕', size=10, color='64748B')
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator='equal', formula=['"완료"'], fill=green_fill, font=green_font),
    )
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator='equal', formula=['"진행중"'], fill=blue_fill, font=blue_font),
    )
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator='equal', formula=['"미착수"'], fill=gray_fill, font=gray_font),
    )

    # 창 고정 (헤더 3행)
    ws.freeze_panes = 'A4'

    return row - 1  # 실제 마지막 데이터 행


def build_summary_sheet(ws, week_label, base_date, members, person_sheet_last_row):
    ws.title = '요약'
    widths = {'A': 4, 'B': 12, 'C': 11, 'D': 10, 'E': 10, 'F': 10, 'G': 13}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # 1행: 제목
    ws.merge_cells('B1:G1')
    ws['B1'] = '인별 과제현황 요약'
    ws['B1'].font = TITLE_FONT
    ws['B1'].fill = TITLE_FILL
    ws['B1'].alignment = CENTER
    ws.row_dimensions[1].height = 28

    # 2행: 부제
    ws.merge_cells('B2:G2')
    ws['B2'] = f"{week_label}   |   기준일 : {base_date.strftime('%Y.%m.%d')}"
    ws['B2'].font = SUB_FONT
    ws['B2'].fill = SUB_FILL
    ws['B2'].alignment = CENTER
    ws.row_dimensions[2].height = 22

    # 4행: 헤더
    headers = ['팀원', '과제 건수', '완료', '진행중', '미착수', '평균 진척율']
    for i, h in enumerate(headers, start=2):  # B열부터
        c = ws.cell(row=4, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER_ALL
    ws.row_dimensions[4].height = 24

    # 인별 시트 참조 범위 (전체 데이터 영역)
    key_range = f"주간업무_인별!$I$4:$I${person_sheet_last_row}"
    status_range = f"주간업무_인별!$H$4:$H${person_sheet_last_row}"
    pct_range = f"주간업무_인별!$G$4:$G${person_sheet_last_row}"
    task_range = f"주간업무_인별!$C$4:$C${person_sheet_last_row}"

    # 5행부터 팀원 데이터
    r = 5
    for name in members:
        # B: 팀원 이름
        b = ws.cell(row=r, column=2, value=name)
        b.font = BODY_FONT
        b.alignment = CENTER
        b.border = BORDER_ALL

        # C: 과제 건수 — 업무내용(C열)이 채워진 행만 카운트 (팀원 KEY = 이름 && C 비어있지 않음)
        # COUNTIFS 로 2조건 처리 : 이름 일치 & 업무내용 비어있지 않음("<>")
        ws.cell(row=r, column=3, value=(
            f'=COUNTIFS({key_range},B{r},{task_range},"<>")'
        ))
        # D: 완료 건수
        ws.cell(row=r, column=4, value=(
            f'=COUNTIFS({key_range},B{r},{task_range},"<>",{status_range},"완료")'
        ))
        # E: 진행중 건수
        ws.cell(row=r, column=5, value=(
            f'=COUNTIFS({key_range},B{r},{task_range},"<>",{status_range},"진행중")'
        ))
        # F: 미착수 건수
        ws.cell(row=r, column=6, value=(
            f'=COUNTIFS({key_range},B{r},{task_range},"<>",{status_range},"미착수")'
        ))
        # G: 평균 진척율 — 진척율이 숫자(값 있음)인 항목만 대상
        #     AVERAGEIFS 에 ">=0" 조건으로 숫자만 필터 (빈 셀은 제외)
        ws.cell(row=r, column=7, value=(
            f'=IFERROR(AVERAGEIFS({pct_range},{key_range},B{r},{pct_range},">=0"),"-")'
        ))

        for col in range(2, 8):
            c = ws.cell(row=r, column=col)
            c.border = BORDER_ALL
            c.font = BODY_FONT
            if col > 2:
                c.alignment = CENTER
        ws.cell(row=r, column=7).number_format = '0.0%'
        ws.row_dimensions[r].height = 22
        r += 1

    # 합계 행
    total_r = r
    ws.cell(row=total_r, column=2, value='합계').font = Font(name='맑은 고딕', size=10, bold=True)
    ws.cell(row=total_r, column=2).alignment = CENTER
    ws.cell(row=total_r, column=2).fill = PatternFill('solid', fgColor='F1F5F9')
    for col in range(3, 7):
        letter = get_column_letter(col)
        ws.cell(row=total_r, column=col, value=f'=SUM({letter}5:{letter}{r - 1})')
        ws.cell(row=total_r, column=col).font = Font(name='맑은 고딕', size=10, bold=True)
        ws.cell(row=total_r, column=col).alignment = CENTER
        ws.cell(row=total_r, column=col).fill = PatternFill('solid', fgColor='F1F5F9')
    # 합계 진척율 = 전체 평균
    ws.cell(row=total_r, column=7, value=(
        f'=IFERROR(AVERAGEIFS({pct_range},{pct_range},">=0"),"-")'
    ))
    ws.cell(row=total_r, column=7).number_format = '0.0%'
    ws.cell(row=total_r, column=7).font = Font(name='맑은 고딕', size=10, bold=True)
    ws.cell(row=total_r, column=7).alignment = CENTER
    ws.cell(row=total_r, column=7).fill = PatternFill('solid', fgColor='F1F5F9')
    for col in range(2, 8):
        ws.cell(row=total_r, column=col).border = BORDER_ALL

    # 하단 주석
    for offset, msg in enumerate([
        '※ 상태 구분 : 완료(진척율 100%) / 진행중(0% 초과 ~ 100% 미만) / 미착수(0%)',
        '※ 평균 진척율은 진척율 기재 과제만 대상',
        '※ 과제 건수는 [주간업무_인별] 시트 업무내용(C열)이 채워진 행 기준',
    ], start=2):
        r2 = total_r + offset
        ws.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=7)
        ws.cell(row=r2, column=2, value=msg).font = FOOT_FONT
        ws.cell(row=r2, column=2).fill = FOOT_FILL
        ws.cell(row=r2, column=2).alignment = LEFT


def build_type_skeleton(ws):
    ws.title = '통합_업무유형별'
    widths = {'A': 22, 'B': 55, 'C': 10, 'D': 12, 'E': 12, 'F': 10, 'G': 10, 'H': 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells('A1:H1')
    ws['A1'] = '계리결산팀 주간업무 현황 (업무유형별 통합)'
    ws['A1'].font = TITLE_FONT
    ws['A1'].fill = TITLE_FILL
    ws['A1'].alignment = CENTER
    ws.row_dimensions[1].height = 28

    ws.merge_cells('A2:H2')
    ws['A2'] = '※ 인별 시트 완성 후 업무유형(Ⅰ~Ⅵ) 기준으로 재분류하여 이곳에 정리'
    ws['A2'].font = SUB_FONT
    ws['A2'].fill = SUB_FILL
    ws['A2'].alignment = CENTER
    ws.row_dimensions[2].height = 22

    headers = ['세부 구분', '업무내용', '담당', '시작일', '완료(예정)', '진척율', '상태', '이슈 · 비고']
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER_ALL
    ws.row_dimensions[3].height = 24

    # 업무유형 6개 카테고리 헤더만 미리 채워둠
    categories = [
        'Ⅰ. IFRS17 결산',
        'Ⅱ. 결산 실무(준비금·비금·계약)',
        'Ⅲ. 차세대 시스템 구축',
        'Ⅳ. 모델·시스템 관리',
        'Ⅴ. 대내외 대응',
        'Ⅵ. 관리회계 및 기타',
    ]
    r = 4
    cat_fill = PatternFill('solid', fgColor='DBEAFE')
    cat_font = Font(name='맑은 고딕', size=11, bold=True, color='1E40AF')
    for cat in categories:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(row=r, column=1, value=cat)
        c.font = cat_font
        c.fill = cat_fill
        c.alignment = LEFT
        c.border = BORDER_ALL
        ws.row_dimensions[r].height = 24
        r += 5  # 카테고리 사이 여백 (4개 빈 행)

    ws.freeze_panes = 'A4'


def main():
    ap = argparse.ArgumentParser(description='계리결산팀 주간업무 인별 템플릿 생성')
    ap.add_argument('--week', required=True, help="주차 라벨 (예: '26년 9월 3주차)")
    ap.add_argument('--base-date', required=True, help='기준일 YYYY-MM-DD')
    ap.add_argument('--out', required=True, help='출력 xlsx 경로')
    ap.add_argument('--rows-per-person', type=int, default=5, help='팀원당 빈 행 수 (기본 5)')
    ap.add_argument('--members', nargs='*', default=None, help='팀원 이름 목록 (미지정 시 기본값 사용)')
    args = ap.parse_args()

    base_date = date.fromisoformat(args.base_date)
    members = args.members or DEFAULT_MEMBERS

    wb = Workbook()
    # 첫 시트 활용
    ws1 = wb.active
    last_row = build_person_sheet(ws1, args.week, base_date, members, args.rows_per_person)

    ws2 = wb.create_sheet('요약')
    build_summary_sheet(ws2, args.week, base_date, members, last_row)

    ws3 = wb.create_sheet('통합_업무유형별')
    build_type_skeleton(ws3)

    out = args.out
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    wb.save(out)
    print(f'OK: {out}')
    print(f'  · 인별 시트 : {len(members)}명 x {args.rows_per_person}행 = {len(members) * args.rows_per_person}행 준비')
    print(f'  · 요약 시트 : 자동집계 (COUNTIFS / AVERAGEIFS)')
    print(f'  · 통합 시트 : 카테고리 6개 스켈레톤 (인별 완성 후 재분류용)')


if __name__ == '__main__':
    main()
